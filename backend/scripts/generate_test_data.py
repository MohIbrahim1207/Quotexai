"""Script to generate realistic test quotation PDF (DMN India) and formatted Excel template."""
from pathlib import Path
import fitz  # PyMuPDF
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
SAMPLE_DIR = ROOT / "sample_data"
TEMPLATE_DIR = ROOT / "templates"
SAMPLE_DIR.mkdir(parents=True, exist_ok=True)
TEMPLATE_DIR.mkdir(parents=True, exist_ok=True)


def generate_dmn_pdf():
    """Generates Quote_41260607.pdf matching real DMN India quotation structure."""
    pdf_path = SAMPLE_DIR / "Quote_41260607.pdf"
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)  # A4

    text = """
DMN INDIA PRIVATE LIMITED
Plot No. 12, Industrial Area, Sector 58, Faridabad - 121004, India
Email: sales@dmn-india.com | Web: www.dmnwestlake.com

QUOTATION / PROFORMA INVOICE

Quote Number: 41260607                          Date: 30/07/2026
Customer: PT. Flow Force Indonesia              Expiry Date: 30/08/2026
Attn: Procurement Dept.                         Currency: EUR (€)
Payment Terms: 100% Advance Against PI          Delivery Terms: Ex-Works

Equipment: ROTARY VALVE DMN BL-300-DAIRY APS
Serial Number: RVNL170001

LINE ITEMS:
----------------------------------------------------------------------------------------------------------------------
Line | Part Number | Description                           | Qty | Unit | Unit Price | Total Price
----------------------------------------------------------------------------------------------------------------------
1    | 23254164    | Lantern ring BL-300                   | 2   | PCS  | 562.00     | 1124.00
2    | 00138737    | Hexagonal bolt M8x20 DIN 933 A2       | 4   | PCS  | 0.32       | 1.28
3    | 23254165    | Cover plate BL-300                    | 2   | PCS  | 28.40      | 56.80
4    | 22210826    | Lip seal 45x65x8 NBR                  | 4   | PCS  | 14.50      | 58.00
5    | 02174072    | Gland packing PTFE-FDA 8x8            | 2   | MTR  | 42.00      | 84.00
6    | 02174084    | O-ring 310x5 FKM 70 Shore A           | 2   | PCS  | 18.20      | 36.40
7    | 02305474    | Rotor BL-300 8-vane AISI 316L         | 1   | PCS  | 1480.00    | 1480.00
8    | 02305314    | Drive shaft BL-300 Hard Chromed       | 1   | PCS  | 620.00     | 620.00
9    | 23072410    | Ball bearing 6209-2RS C3              | 2   | PCS  | 35.00      | 70.00
10   | 23084110    | Circlip D85 DIN 472                   | 2   | PCS  | 1.80       | 3.60
11   | 22171642    | Grease nipple G1/8 straight           | 4   | PCS  | 2.10       | 8.40
----------------------------------------------------------------------------------------------------------------------

Packaging & Forwarding: EUR 0.00
Miscellaneous Charges: EUR 0.00
Lines Total: EUR 3542.48
Grand Total: EUR 3542.48

Remarks:
1. Part numbers starting with leading zeros (e.g. 00138737, 02174072, 02174084, 02305474, 02305314) are exact OEM codes.
2. Prices are exclusive of import duties and local taxes.
3. Country of Origin: Netherlands / India.
"""
    # Insert text into page
    rect = fitz.Rect(36, 36, 559, 806)
    page.insert_textbox(rect, text, fontsize=9, fontname="helv", color=(0.1, 0.1, 0.15))
    doc.save(str(pdf_path))
    doc.close()
    print(f"Generated sample PDF at {pdf_path}")
    return pdf_path


def generate_excel_template():
    """Generates ENQ-2026-07-2549.xlsx template with merged headers, styling, and formulas."""
    template_path = TEMPLATE_DIR / "ENQ-2026-07-2549.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Quotation"
    ws.views.sheetView[0].showGridLines = True

    # Styling definitions
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")  # Navy Blue
    title_font = Font(name="Segoe UI", size=16, bold=True, color="1E3A8A")
    th_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    bold_font = Font(name="Segoe UI", size=10, bold=True)
    regular_font = Font(name="Segoe UI", size=10)
    
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    
    header_border = Border(
        left=Side(style='thin', color='1E3A8A'),
        right=Side(style='thin', color='1E3A8A'),
        top=Side(style='medium', color='1E3A8A'),
        bottom=Side(style='medium', color='1E3A8A')
    )

    # Title Banner
    ws.merge_cells("A2:G2")
    ws["A2"] = "INDUSTRIAL EQUIPMENT & SPARES QUOTATION"
    ws["A2"].font = title_font
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 30

    # Company & Reference Info
    ws["A4"] = "Company:"
    ws["A4"].font = bold_font
    ws["B4"] = "FLOW FORCE INDUSTRIAL SOLUTIONS"
    ws["B4"].font = regular_font

    ws["A5"] = "Customer:"
    ws["A5"].font = bold_font
    ws["B5"] = "PT. Flow Force Indonesia"  # Cell B5
    ws["B5"].font = regular_font

    ws["F4"] = "Enquiry No:"
    ws["F4"].font = bold_font
    ws["G4"] = "ENQ-2026-07-2549"
    ws["G4"].font = bold_font

    ws["F5"] = "Quote Ref:"
    ws["F5"].font = bold_font
    ws["G5"] = "41260607"  # Cell G5
    ws["G5"].font = regular_font

    ws["F6"] = "Date:"
    ws["F6"].font = bold_font
    ws["G6"] = "30/07/2026"
    ws["G6"].font = regular_font

    ws["F7"] = "Currency:"
    ws["F7"].font = bold_font
    ws["G7"] = "EUR"
    ws["G7"].font = regular_font

    # Table Headers at Row 10
    headers = [
        ("A", "Item No", 8),
        ("B", "Part Number", 18),
        ("C", "Description", 38),
        ("D", "Qty", 8),
        ("E", "UOM", 8),
        ("F", "Unit Price (€)", 16),
        ("G", "Total Price (€)", 18),
    ]

    header_row = 10
    ws.row_dimensions[header_row].height = 24

    for col_letter, title, width in headers:
        cell = ws[f"{col_letter}{header_row}"]
        cell.value = title
        cell.font = th_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = header_border
        ws.column_dimensions[col_letter].width = width

    # Add 12 formatted placeholder rows (11 to 22)
    for r in range(11, 23):
        ws.row_dimensions[r].height = 20
        ws[f"A{r}"].alignment = Alignment(horizontal="center")
        ws[f"B{r}"].alignment = Alignment(horizontal="center")
        ws[f"B{r}"].number_format = "@"  # Crucial: Text format for part numbers
        ws[f"C{r}"].alignment = Alignment(horizontal="left")
        ws[f"D{r}"].alignment = Alignment(horizontal="right")
        ws[f"D{r}"].number_format = "#,##0.00"
        ws[f"E{r}"].alignment = Alignment(horizontal="center")
        ws[f"F{r}"].alignment = Alignment(horizontal="right")
        ws[f"F{r}"].number_format = "#,##0.00"
        ws[f"G{r}"].alignment = Alignment(horizontal="right")
        ws[f"G{r}"].number_format = "#,##0.00"
        ws[f"G{r}"].value = f"=D{r}*F{r}"

        for col_letter, _, _ in headers:
            ws[f"{col_letter}{r}"].border = thin_border
            ws[f"{col_letter}{r}"].font = regular_font

    # Totals Section at Row 23
    total_row = 23
    ws.merge_cells(f"A{total_row}:F{total_row}")
    ws[f"A{total_row}"] = "TOTAL AMOUNT (EUR)"
    ws[f"A{total_row}"].font = bold_font
    ws[f"A{total_row}"].alignment = Alignment(horizontal="right", vertical="center")
    
    ws[f"G{total_row}"] = f"=SUM(G11:G22)"
    ws[f"G{total_row}"].font = bold_font
    ws[f"G{total_row}"].alignment = Alignment(horizontal="right", vertical="center")
    ws[f"G{total_row}"].number_format = "#,##0.00"
    ws[f"G{total_row}"].border = Border(top=Side(style='thin', color='1E3A8A'), bottom=Side(style='double', color='1E3A8A'))

    wb.save(str(template_path))
    print(f"Generated sample Excel template at {template_path}")
    return template_path


if __name__ == "__main__":
    generate_dmn_pdf()
    generate_excel_template()
