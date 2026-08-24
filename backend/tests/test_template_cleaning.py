"""Automated tests to verify that old/sample template items NEVER survive in the generated Excel."""
import pytest
from pathlib import Path
from app.schemas.quotation import QuoteItem, QuotationData
from app.services.excel_service import excel_service
from app.config import TEMPLATE_DIR, OUTPUT_DIR
import openpyxl


def test_old_sample_part_numbers_are_completely_removed():
    """Verify that old sample part numbers (03431814, 22462055, 02176160, etc.) are 100% removed."""
    sample_template_parts = [
        "03431814",
        "22462055",
        "02176160",
        "02175055",
        "46016050",
        "00354420",
        "00302550",
        "00185452",
    ]

    # Provide only 3 fresh items
    pdf_items = [
        QuoteItem(line_number=1, part_number="NEW-001", description="Fresh Part 1", quantity=5.0, unit_price=100.0),
        QuoteItem(line_number=2, part_number="NEW-002", description="Fresh Part 2", quantity=2.0, unit_price=250.0),
        QuoteItem(line_number=3, part_number="NEW-003", description="Fresh Part 3", quantity=1.0, unit_price=50.0),
    ]

    quote = QuotationData(
        quote_number="CLEAN-TEST-001",
        customer="Testing Client",
        currency="EUR",
        items=pdf_items,
    )

    # Test with root template if available
    root_template = Path(__file__).resolve().parents[2] / "templates" / "ENQ-2026-07-2549.xlsx"
    backend_template = TEMPLATE_DIR / "ENQ-2026-07-2549.xlsx"
    template_to_use = root_template if root_template.exists() else backend_template

    out_path = OUTPUT_DIR / "test_sample_removal_output.xlsx"
    count = excel_service.generate_quotation_excel(
        template_path=template_to_use,
        output_path=out_path,
        quotation=quote,
    )

    assert count == 3

    # Inspect all cells in output workbook
    wb = openpyxl.load_workbook(str(out_path), data_only=False)
    sheet = wb.active

    all_cell_values = []
    for r in range(1, sheet.max_row + 1):
        for c in range(1, sheet.max_column + 1):
            val = sheet.cell(row=r, column=c).value
            if val is not None:
                all_cell_values.append(str(val))

    # None of the old sample part numbers must appear
    for old_pn in sample_template_parts:
        assert old_pn not in all_cell_values, f"Old template sample part {old_pn} survived in generated Excel!"

    # The 3 PDF items must appear
    assert "NEW-001" in all_cell_values
    assert "NEW-002" in all_cell_values
    assert "NEW-003" in all_cell_values

    wb.close()
