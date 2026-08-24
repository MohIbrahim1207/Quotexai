"""Automated tests for multiple template families (5-column ENQ vs 7-column standard) ensuring no column shifts."""
import pytest
from pathlib import Path
import openpyxl
from app.services.excel_service import excel_service
from app.services.pdf_service import pdf_service
from app.services.ai_service import DeterministicFallbackProvider
from app.services.quotation_service import quotation_orchestrator
from app.schemas.quotation import GenerateExcelRequest, QuoteItem, QuotationData
from app.config import TEMPLATE_DIR, OUTPUT_DIR, SAMPLE_DIR


def test_5_column_enq_template_physical_placement():
    """Verify that 5-column ENQ template (A:No, B:Part Number, C:Description, D:Unit Price, E:Qty) has exact values in D and E, and F/G empty."""
    pdf_path = SAMPLE_DIR / "Quote_Form_41260607.pdf"
    if not pdf_path.exists():
        pdf_path = Path(__file__).resolve().parents[2] / "sample_data" / "Quote_Form_41260607.pdf"
    template_path = TEMPLATE_DIR / "ENQ-2026-07-2549.xlsx"
    out_path = OUTPUT_DIR / "test_5_col_physical_verify.xlsx"

    pdf = pdf_service.extract_content(pdf_path)
    quote = DeterministicFallbackProvider().extract_quotation_sync(pdf.full_text, filename="Quote_Form_41260607.pdf")

    # Generate via orchestrator (full pipeline)
    req = GenerateExcelRequest(
        quotation=quote,
        template_id="ENQ-2026-07-2549.xlsx",
        pricing_mode="quoted_price",
        margin_percent=0.0,
        supplier_discount_percent=0.0,
    )
    res = quotation_orchestrator.generate_excel_quotation(req)
    gen_file = OUTPUT_DIR / res.file_id

    wb = openpyxl.load_workbook(str(gen_file), data_only=True)
    ws = wb.active

    # Row 6: Equipment Section 1 Header
    assert ws["A6"].value == "ROTARY VALVE DMN BL-300-DAIRY APS (S/N: RVNL170001)"
    # Row 7: Item 1 (23254164)
    assert ws["A7"].value == 1
    assert ws["B7"].value == "23254164"
    assert ws["D7"].value == 562.00, f"Expected Unit Price 562.00 in D7, got {ws['D7'].value}"
    assert ws["E7"].value == 2, f"Expected Qty 2 in E7, got {ws['E7'].value}"
    assert ws["F7"].value in (None, ""), f"F7 must be empty, got {ws['F7'].value}"
    assert ws["G7"].value in (None, ""), f"G7 must be empty, got {ws['G7'].value}"

    # Row 8: Item 2 (00138737)
    assert ws["A8"].value == 2
    assert ws["B8"].value == "00138737"
    assert ws["D8"].value == 0.32, f"Expected Unit Price 0.32 in D8, got {ws['D8'].value}"
    assert ws["E8"].value == 4, f"Expected Qty 4 in E8, got {ws['E8'].value}"
    assert ws["F8"].value in (None, ""), f"F8 must be empty, got {ws['F8'].value}"
    assert ws["G8"].value in (None, ""), f"G8 must be empty, got {ws['G8'].value}"

    # Confirm no NOS in E
    for r in range(6, 20):
        assert ws[f"E{r}"].value != "NOS", f"E{r} contains 'NOS'!"

    wb.close()


def test_7_column_standard_template_physical_placement():
    """Verify that 7-column template (A:Item No, B:Part Number, C:Description, D:Qty, E:UOM, F:Unit Price, G:Total Price) writes to all 7 columns correctly."""
    # Create 7-column template
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Standard7Col"
    ws["A10"] = "Item No"
    ws["B10"] = "Part Number"
    ws["C10"] = "Description"
    ws["D10"] = "Qty"
    ws["E10"] = "UOM"
    ws["F10"] = "Unit Price"
    ws["G10"] = "Total Price"

    tmpl_path = OUTPUT_DIR / "temp_7_col_template.xlsx"
    wb.save(str(tmpl_path))
    wb.close()

    items = [
        QuoteItem(
            line_number=1,
            part_number="23254164",
            description="Lantern ring",
            quantity=2.0,
            unit="NOS",
            unit_price=562.00,
            total_price=1124.00,
        ),
        QuoteItem(
            line_number=2,
            part_number="00138737",
            description="Hex bolt",
            quantity=4.0,
            unit="NOS",
            unit_price=0.32,
            total_price=1.28,
        ),
    ]
    quote = QuotationData(quote_number="7COL-01", items=items)

    out_path = OUTPUT_DIR / "test_7_col_out.xlsx"
    count = excel_service.generate_quotation_excel(tmpl_path, out_path, quote)
    assert count == 2

    wb_out = openpyxl.load_workbook(str(out_path), data_only=False)
    ws_out = wb_out.active

    # Row 11: Item 1
    assert ws_out["A11"].value == 1
    assert ws_out["B11"].value == "23254164"
    assert ws_out["C11"].value == "Lantern ring"
    assert ws_out["D11"].value == 2.0
    assert ws_out["E11"].value == "NOS"
    assert ws_out["F11"].value == 562.00
    assert ws_out["G11"].value == 1124.00 or ws_out["G11"].value == "=D11*F11"

    # Row 12: Item 2
    assert ws_out["A12"].value == 2
    assert ws_out["B12"].value == "00138737"
    assert ws_out["D12"].value == 4.0
    assert ws_out["E12"].value == "NOS"
    assert ws_out["F12"].value == 0.32

    wb_out.close()
