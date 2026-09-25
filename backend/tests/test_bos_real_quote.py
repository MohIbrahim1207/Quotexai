"""Regression tests for the real BOS Homogenisers quotation (H3-400P Flow Force Sept.pdf)."""
import pytest
from pathlib import Path
import openpyxl

from app.config import SAMPLE_DIR, TEMPLATE_DIR, OUTPUT_DIR
from app.services.pdf_service import pdf_service
from app.services.ai_service import DeterministicFallbackProvider
from app.services.validation_service import validation_service
from app.services.quotation_service import quotation_orchestrator
from app.schemas.quotation import GenerateExcelRequest


def test_bos_single_main_item_and_metadata():
    """Verify that generic descriptive parser correctly extracts BOS Homogeniser quote as 1 main item with no phantom options."""
    pdf_path = SAMPLE_DIR / "H3-400P_Flow_Force_Sept.pdf"
    if not pdf_path.exists():
        pdf_path = Path(__file__).resolve().parents[1] / "sample_data" / "H3-400P_Flow_Force_Sept.pdf"
    if not pdf_path.exists():
        pdf_path = Path(__file__).resolve().parents[2] / "uploads" / "c09106e05c3a_H3-400P_Flow_Force_Sept.pdf"
    assert pdf_path.exists(), f"PDF not found at {pdf_path}"

    extraction = pdf_service.extract_content(pdf_path)
    parser = DeterministicFallbackProvider()
    quotation = parser.extract_quotation_sync(extraction.full_text, filename="H3-400P_Flow_Force_Sept.pdf")

    # 1. Quotation Layout & Metadata
    assert quotation.layout_type == "descriptive"
    assert quotation.supplier_name == "Bos Homogenisers B.V."
    assert quotation.customer == "Flow Force"
    assert quotation.quotation_date == "2026-09-21"
    assert quotation.subject == "H3-400P"

    # 2. Exactly 1 Main Quoted Item (Catalog options are NOT extracted as purchased items)
    assert len(quotation.items) == 1, f"Expected exactly 1 main item, got {len(quotation.items)}"
    item = quotation.items[0]

    # 3. Main Quoted Machine Fields
    assert item.description == "BOS high-pressure homogeniser H-Series H3-400P"
    assert item.quantity == 1.0
    assert item.unit == "NOS"
    assert item.currency == "EUR"
    assert item.unit_price == 26268.00
    assert item.discount_percent == 15.0
    assert item.net_price == 22327.80
    assert item.total_price == 22327.80

    # 4. Fields that do not exist must remain blank / null
    assert item.part_number in ("", None)
    assert item.commodity_code is None

    # 5. Catalog option codes must NEVER be extracted as purchased line items
    forbidden_option_codes = ["HEC14", "HEC09", "HEC13", "HCL01", "HCL18", "HPR01", "HAU06", "HEC15"]
    for opt_code in forbidden_option_codes:
        assert all(opt_code not in (it.part_number or "") for it in quotation.items)
        assert all(it.description != opt_code for it in quotation.items)

    # 6. Validation
    val_quote, val_summary = validation_service.validate_quotation(quotation, extraction.full_text)
    assert val_summary.is_valid is True
    assert val_summary.errors_count == 0


def test_bos_physical_excel_cells():
    """Verify generated Excel physical cells reopen with openpyxl: A7=1, B7=blank, C7=desc, D7=26268, E7=1."""
    pdf_path = SAMPLE_DIR / "H3-400P_Flow_Force_Sept.pdf"
    if not pdf_path.exists():
        pdf_path = Path(__file__).resolve().parents[1] / "sample_data" / "H3-400P_Flow_Force_Sept.pdf"
    if not pdf_path.exists():
        pdf_path = Path(__file__).resolve().parents[2] / "uploads" / "c09106e05c3a_H3-400P_Flow_Force_Sept.pdf"

    extraction = pdf_service.extract_content(pdf_path)
    quote = DeterministicFallbackProvider().extract_quotation_sync(extraction.full_text, filename="H3-400P_Flow_Force_Sept.pdf")

    req = GenerateExcelRequest(
        quotation=quote,
        template_id="ENQ-2026-07-2549.xlsx",
        pricing_mode="quoted_price",
        margin_percent=0.0,
        supplier_discount_percent=0.0,
    )
    res = quotation_orchestrator.generate_excel_quotation(req)
    gen_file = OUTPUT_DIR / res.file_id
    assert gen_file.exists(), f"Excel file {gen_file} was not generated."

    wb = openpyxl.load_workbook(str(gen_file), data_only=True)
    ws = wb.active

    # Row 6: Equipment Section Header
    assert ws["A6"].value == "BOS high-pressure homogeniser H-Series H3-400P"

    # Row 7: Physical Cells Verification
    assert ws["A7"].value == 1, f"Expected A7=1, got {ws['A7'].value}"
    assert ws["B7"].value in (None, ""), f"Expected B7 to be blank, got {ws['B7'].value}"
    assert ws["C7"].value == "BOS high-pressure homogeniser H-Series H3-400P", f"Expected C7 description, got {ws['C7'].value}"
    assert ws["D7"].value == 26268.00, f"Expected D7=26268.00, got {ws['D7'].value}"
    assert ws["E7"].value == 1, f"Expected E7=1, got {ws['E7'].value}"

    wb.close()
