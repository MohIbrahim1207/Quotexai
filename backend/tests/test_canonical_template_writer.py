"""Automated test matching the user's canonical header detection and template filling specification."""
import pytest
from pathlib import Path
import openpyxl
from app.services.excel_service import (
    HEADER_SYNONYMS,
    detect_column_mapping,
    find_first_data_row,
    excel_service,
)
from app.services.pdf_service import pdf_service
from app.services.ai_service import DeterministicFallbackProvider
from app.config import TEMPLATE_DIR, OUTPUT_DIR, SAMPLE_DIR


def test_canonical_header_synonyms_and_detection():
    """Verify detect_column_mapping and find_first_data_row on ENQ-2026-07-2549.xlsx."""
    template_path = TEMPLATE_DIR / "ENQ-2026-07-2549.xlsx"
    wb = openpyxl.load_workbook(str(template_path))
    ws = wb.active

    mapping, last_header_row, _ = detect_column_mapping(ws)
    assert mapping == {
        "line_number": "A",
        "part_number": "B",
        "description": "C",
        "unit_price": "D",
        "quantity": "E",
    }
    assert last_header_row == 4

    start_row = find_first_data_row(ws, mapping, last_header_row)
    assert start_row == 6
    wb.close()


def test_canonical_user_acceptance_quote_41260607_physical_cells():
    """Run full physical cell inspection on Quote 41260607 filled into ENQ-2026-07-2549.xlsx."""
    pdf_path = SAMPLE_DIR / "Quote_Form_41260607.pdf"
    if not pdf_path.exists():
        pdf_path = Path(__file__).resolve().parents[2] / "sample_data" / "Quote_Form_41260607.pdf"
    template_path = TEMPLATE_DIR / "ENQ-2026-07-2549.xlsx"
    out_path = OUTPUT_DIR / "canonical_acceptance_test_out.xlsx"

    pdf = pdf_service.extract_content(pdf_path)
    quote = DeterministicFallbackProvider().extract_quotation_sync(pdf.full_text, filename="Quote_Form_41260607.pdf")

    count = excel_service.generate_quotation_excel(template_path, out_path, quote)
    assert count == 11

    # Physical cell inspection
    expected = {
        "A6": "ROTARY VALVE DMN BL-300-DAIRY APS (S/N: RVNL170001)",
        "A7": 1, "B7": "23254164", "D7": 562.00, "E7": 2,
        "A8": 2, "B8": "00138737", "D8": 0.32, "E8": 4,
        "F7": None, "G7": None,
    }

    wb = openpyxl.load_workbook(str(out_path), data_only=True)
    ws = wb.active

    for coord, exp in expected.items():
        act = ws[coord].value
        assert act == exp, f"Physical cell {coord} mismatch: expected {exp!r}, got {act!r}"

    # Confirm bad pattern is clean
    for r in range(6, 20):
        assert ws[f"E{r}"].value != "NOS", f"E{r} contains 'NOS'!"
        assert ws[f"F{r}"].value in (None, ""), f"F{r} contains unexpected value {ws[f'F{r}'].value!r}"
        assert ws[f"G{r}"].value in (None, ""), f"G{r} contains unexpected value {ws[f'G{r}'].value!r}"

    wb.close()
