"""Test case for non-inventory charges (e.g. S&H on quote 11009223) ensuring no false-positive missing part number errors."""
import pytest
from pathlib import Path
import openpyxl
from app.services.pdf_service import pdf_service
from app.services.ai_service import DeterministicFallbackProvider
from app.services.validation_service import validation_service, is_non_inventory_charge
from app.services.excel_service import excel_service
from app.schemas.quotation import QuotationData, QuoteItem
from app.config import SAMPLE_DIR, OUTPUT_DIR, TEMPLATE_DIR


def test_is_non_inventory_charge_detection():
    """Verify is_non_inventory_charge correctly classifies charge lines vs real parts."""
    charge_item1 = QuoteItem(
        line_number=1,
        part_number="S&H",
        description="SHIPPING & HANDLING",
        quantity=1.0,
        unit_price=1200.0,
    )
    assert is_non_inventory_charge(charge_item1) is True

    charge_item2 = QuoteItem(
        line_number=2,
        part_number="",
        description="FREIGHT AND INSURANCE CHARGE",
        quantity=1.0,
        unit_price=500.0,
    )
    assert is_non_inventory_charge(charge_item2) is True

    real_part_item = QuoteItem(
        line_number=3,
        part_number="13120051",
        description="BLADE - DS 225 .300 420",
        quantity=16.0,
        unit_price=416.25,
    )
    assert is_non_inventory_charge(real_part_item) is False

    missing_part_item = QuoteItem(
        line_number=4,
        part_number="",
        description="HEXAGONAL NUT M8 STAINLESS",
        quantity=10.0,
        unit_price=2.50,
    )
    assert is_non_inventory_charge(missing_part_item) is False


def test_validation_passes_for_fitzpatrick_quote_with_sh_charge():
    """Verify validation produces 0 critical errors on Fitzpatrick quote 11009223 containing an S&H line."""
    pdf_path = SAMPLE_DIR / "Fitzpatrick_11009223_PT_Flow_Force_Indonesia_.pdf"
    if not pdf_path.exists():
        pdf_path = Path(__file__).resolve().parents[2] / "sample_data" / "Fitzpatrick_11009223_PT_Flow_Force_Indonesia_.pdf"

    ext = pdf_service.extract_content(pdf_path)
    quote = DeterministicFallbackProvider().extract_quotation_sync(ext.full_text, filename="Fitzpatrick_11009223.pdf")

    assert quote.quote_number == "11009223"
    assert len(quote.items) == 2

    # Line 1: Part
    assert quote.items[0].part_number == "13120051"
    assert quote.items[0].quantity == 16.0
    assert quote.items[0].unit_price == 416.25

    # Line 2: S&H Charge
    assert quote.items[1].part_number == "S&H"
    assert quote.items[1].quantity == 1.0
    assert quote.items[1].unit_price == 1200.0

    # Validate
    val_quote, summary = validation_service.validate_quotation(quote, ext.full_text)
    assert summary.errors_count == 0, f"Expected 0 critical errors, got {summary.errors_count}"
    assert summary.is_valid is True

    # Info issue exists for S&H
    info_issues = [i for i in summary.issues if i.issue_type == "info"]
    assert len(info_issues) >= 1
    assert "non-inventory charge" in info_issues[0].message.lower()


def test_validation_still_fails_for_genuinely_missing_part_number_on_real_part():
    """Verify that a physical part missing its part number is still flagged as a critical error."""
    items = [
        QuoteItem(
            line_number=1,
            part_number="13120051",
            description="BLADE - DS 225 .300 420",
            quantity=16.0,
            unit_price=416.25,
        ),
        QuoteItem(
            line_number=2,
            part_number="",
            description="ROTOR SEAL RING 300MM",
            quantity=2.0,
            unit_price=150.00,
        ),
    ]
    quote = QuotationData(quote_number="TEST-MISSING-PN", items=items)
    val_quote, summary = validation_service.validate_quotation(quote, "")

    assert summary.errors_count == 1
    assert summary.is_valid is False
    assert any("Missing part number" in i.message for i in summary.issues if i.issue_type == "error")


def test_fitzpatrick_excel_generation_includes_both_items():
    """Verify that Excel generation for Fitzpatrick quote includes both the blade and the S&H line."""
    pdf_path = SAMPLE_DIR / "Fitzpatrick_11009223_PT_Flow_Force_Indonesia_.pdf"
    if not pdf_path.exists():
        pdf_path = Path(__file__).resolve().parents[2] / "sample_data" / "Fitzpatrick_11009223_PT_Flow_Force_Indonesia_.pdf"
    template_path = TEMPLATE_DIR / "ENQ-2026-07-2549.xlsx"
    out_path = OUTPUT_DIR / "fitzpatrick_11009223_out.xlsx"

    ext = pdf_service.extract_content(pdf_path)
    quote = DeterministicFallbackProvider().extract_quotation_sync(ext.full_text, filename="Fitzpatrick_11009223.pdf")

    count = excel_service.generate_quotation_excel(template_path, out_path, quote)
    assert count == 2

    wb = openpyxl.load_workbook(str(out_path), data_only=True)
    ws = wb.active

    # Row 6: Item 1 (13120051)
    assert ws["A6"].value == 1
    assert ws["B6"].value == "13120051"
    assert ws["D6"].value == 416.25
    assert ws["E6"].value == 16.0
    assert ws["F6"].value in (None, "")

    # Row 7: Item 2 (S&H)
    assert ws["A7"].value == 2
    assert ws["B7"].value == "S&H"
    assert ws["D7"].value == 1200.00
    assert ws["E7"].value == 1.0
    assert ws["F7"].value in (None, "")

    wb.close()
