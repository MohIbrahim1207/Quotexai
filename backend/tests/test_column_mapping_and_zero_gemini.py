"""Tests for deterministic Excel column mapping, non-existent field handling, and zero Gemini API calls during Excel operations."""
import pytest
from pathlib import Path
import openpyxl
from app.schemas.quotation import QuoteItem, QuotationData, EquipmentGroup
from app.services.excel_service import excel_service
from app.services.ai_service import GeminiProvider
from app.config import TEMPLATE_DIR, OUTPUT_DIR


def test_custom_column_order_mapping_without_gemini(monkeypatch):
    """Verify that template analysis and Excel writing correctly follow whatever column order is defined in the template."""
    # Ensure Gemini is NEVER called
    def fail_gemini(*args, **kwargs):
        raise AssertionError("Gemini was invoked during Excel generation/mapping! Rule violation.")
    monkeypatch.setattr(GeminiProvider, "extract_quotation", fail_gemini)

    # 1. Create a dynamic template with non-standard column order:
    # A: No. | B: Description | C: Qty | D: Part Number | E: Unit Price
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "CustomOrder"
    sheet["A4"] = "No."
    sheet["B4"] = "Description"
    sheet["C4"] = "Qty"
    sheet["D4"] = "Part Number"
    sheet["E4"] = "Unit Price"

    custom_tmpl_path = OUTPUT_DIR / "temp_custom_order_template.xlsx"
    wb.save(str(custom_tmpl_path))
    wb.close()

    # 2. Analyze template
    analysis = excel_service.analyze_template(custom_tmpl_path)
    mapping = analysis.suggested_mapping

    assert mapping.line_number_column == "A"
    assert mapping.description_column == "B"
    assert mapping.quantity_column == "C"
    assert mapping.part_number_column == "D"
    assert mapping.unit_price_column == "E"

    # 3. Generate Excel
    quote = QuotationData(
        quote_number="CUSTOM-ORDER-01",
        items=[
            QuoteItem(line_number=1, part_number="00138737", description="Hex Bolt", quantity=4.0, unit_price=0.32),
            QuoteItem(line_number=2, part_number="23254164", description="Lantern Ring", quantity=2.0, unit_price=562.00),
        ],
    )

    out_path = OUTPUT_DIR / "test_custom_order_out.xlsx"
    count = excel_service.generate_quotation_excel(
        template_path=custom_tmpl_path,
        output_path=out_path,
        quotation=quote,
    )
    assert count == 2

    # 4. Inspect generated Excel
    wb_out = openpyxl.load_workbook(str(out_path), data_only=False)
    sh_out = wb_out.active

    # Row 5 (start_row): Line 1
    assert sh_out["A5"].value == 1
    assert sh_out["B5"].value == "Hex Bolt"
    assert sh_out["C5"].value == 4.0
    assert sh_out["D5"].value == "00138737"
    assert sh_out["D5"].number_format == "@"
    assert sh_out["E5"].value == 0.32

    # Row 6: Line 2
    assert sh_out["A6"].value == 2
    assert sh_out["B6"].value == "Lantern Ring"
    assert sh_out["C6"].value == 2.0
    assert sh_out["D6"].value == "23254164"
    assert sh_out["D6"].number_format == "@"
    assert sh_out["E6"].value == 562.00

    wb_out.close()


def test_unmapped_fields_are_not_written(monkeypatch):
    """Verify that if the template does not have columns for Unit, Discount, Total Price, Currency, or Commodity Code, they are NOT written into other columns."""
    def fail_gemini(*args, **kwargs):
        raise AssertionError("Gemini was invoked during Excel generation! Rule violation.")
    monkeypatch.setattr(GeminiProvider, "extract_quotation", fail_gemini)

    tmpl_path = TEMPLATE_DIR / "ENQ-2026-07-2549.xlsx"
    analysis = excel_service.analyze_template(tmpl_path)
    mapping = analysis.suggested_mapping

    # In ENQ-2026-07-2549.xlsx, there is NO unit column, discount column, or commodity code column
    assert mapping.unit_column is None
    assert mapping.discount_column is None

    quote = QuotationData(
        quote_number="NO-UNMAPPED-COLS",
        items=[
            QuoteItem(
                line_number=1,
                part_number="02174072",
                description="O-ring",
                quantity=2.0,
                unit="EA",
                unit_price=18.40,
                discount_percent=30.0,
                total_price=25.76,
                commodity_code="84819090",
            )
        ]
    )

    out_path = OUTPUT_DIR / "test_unmapped_cols_out.xlsx"
    excel_service.generate_quotation_excel(
        template_path=tmpl_path,
        output_path=out_path,
        quotation=quote,
    )

    wb = openpyxl.load_workbook(str(out_path), data_only=False)
    sheet = wb.active

    # Check Row 6 (start_row) or Row 7
    # Part number in B, Unit price in D (18.4), Qty in E (2.0)
    # Ensure EA, 30.0%, 84819090 are NOT written in D or E
    found_item = False
    for r in range(6, 12):
        if sheet[f"B{r}"].value == "02174072":
            found_item = True
            assert sheet[f"D{r}"].value == 18.40, f"Expected 18.40 in Unit Price, got {sheet[f'D{r}'].value}"
            assert sheet[f"E{r}"].value == 2.0, f"Expected 2.0 in Qty, got {sheet[f'E{r}'].value}"
            # Ensure Commodity Code is not in C or F
            assert sheet[f"C{r}"].value == "O-ring"
            assert sheet[f"F{r}"].value is None

    wb.close()
    assert found_item, "Item 02174072 not found in expected row"
